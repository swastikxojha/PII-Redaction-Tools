from __future__ import annotations

import io
import logging
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from app.detectors.base import DetectionContext
from app.processors.run_processor import apply_replacements_to_paragraph
from app.processors.image_processor import ImageProcessor


@dataclass
class ProcessingStats:
    paragraphs: int=0
    tables: int=0
    table_cells: int=0
    headers: int=0
    footers: int=0
    embedded_images: int=0
    ocr_images: int=0
    redacted_images: int=0
    entities: list=field(default_factory=list)
    warnings: list[str]=field(default_factory=list)

    def add_entities(self, entities):
        self.entities.extend(entities)

    @property
    def counts(self):
        return dict(Counter(e.type.value for e in self.entities))


class DocxProcessor:
    def __init__(self, text_processor, pseudonymizer, logger=None):
        self.text_processor=text_processor
        self.pseudonymizer=pseudonymizer
        self.image_processor=ImageProcessor(text_processor)
        self.logger=logger or logging.getLogger(__name__)

    def process(self, input_path: str | Path, output_path: str | Path) -> ProcessingStats:
        input_path=Path(input_path); output_path=Path(output_path)
        output_path.parent.mkdir(parents=True,exist_ok=True)
        doc=Document(str(input_path))
        stats=ProcessingStats()

        # Prime document-wide catalogs (especially person names) before modifying any text.
        prime_texts=[]
        prime_texts.extend([p.text for p in doc.paragraphs if p.text])
        for t in doc.tables:
            for row in t.rows:
                for cell in row.cells:
                    prime_texts.extend([p.text for p in cell.paragraphs if p.text])
        for section in doc.sections:
            prime_texts.extend([p.text for p in section.header.paragraphs if p.text])
            prime_texts.extend([p.text for p in section.footer.paragraphs if p.text])
        self.text_processor.prime(prime_texts)

        # Body paragraphs
        for i,p in enumerate(doc.paragraphs):
            stats.paragraphs += 1
            self._process_paragraph(p, DetectionContext(location="body", block_id=f"p{i}"), stats)

        # Tables
        stats.tables=len(doc.tables)
        for ti,t in enumerate(doc.tables):
            for ri,row in enumerate(t.rows):
                for ci,cell in enumerate(row.cells):
                    stats.table_cells += 1
                    for pi,p in enumerate(cell.paragraphs):
                        self._process_paragraph(p, DetectionContext(location="table", block_id=f"t{ti}r{ri}c{ci}p{pi}"), stats)

        # Headers and footers, per section.
        for si,s in enumerate(doc.sections):
            for part_name,container in (("header",s.header),("footer",s.footer)):
                for pi,p in enumerate(container.paragraphs):
                    if part_name=="header": stats.headers += 1
                    else: stats.footers += 1
                    self._process_paragraph(p, DetectionContext(location=part_name, block_id=f"{part_name}{si}:{pi}"), stats)
                for ti,t in enumerate(container.tables):
                    for ri,row in enumerate(t.rows):
                        for ci,cell in enumerate(row.cells):
                            for pi,p in enumerate(cell.paragraphs):
                                if part_name=="header": stats.headers += 1
                                else: stats.footers += 1
                                self._process_paragraph(p, DetectionContext(location=part_name, block_id=f"{part_name}{si}:t{ti}r{ri}c{ci}p{pi}"), stats)

        # python-docx saves XML; afterwards replace media parts directly so floating/anchored images remain positioned.
        temp=output_path.with_suffix(".text.docx")
        doc.save(str(temp))
        self._process_embedded_images(temp, output_path, stats)
        temp.unlink(missing_ok=True)
        return stats

    def _process_paragraph(self, paragraph, context, stats):
        text=paragraph.text or ""
        if not text.strip():
            return
        entities=self.text_processor.detect(text, context)
        if not entities:
            return
        stats.add_entities(entities)
        replacements=[(e,self.pseudonymizer.replacement_for(e)) for e in entities]
        apply_replacements_to_paragraph(paragraph, replacements)

    def _process_embedded_images(self, docx_path: Path, output_path: Path, stats: ProcessingStats):
        buf=io.BytesIO()
        with zipfile.ZipFile(docx_path, "r") as zin:
            media_names=[n for n in zin.namelist() if n.startswith("word/media/")]
            stats.embedded_images=len(media_names)
            replaced={}
            for name in media_names:
                raw=zin.read(name)
                result, out=self.image_processor.analyze_bytes(name, raw)
                if result.ocr_text:
                    stats.ocr_images += 1
                if result.sensitive:
                    stats.redacted_images += 1
                    replaced[name]=out
                    self.logger.info("Redacted embedded image: %s", name)
            with zipfile.ZipFile(output_path, "w", compression=zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data=replaced.get(item.filename, zin.read(item.filename))
                    zout.writestr(item, data)
