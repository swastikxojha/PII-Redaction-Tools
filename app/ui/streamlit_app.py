from __future__ import annotations

from pathlib import Path
import tempfile

try:
    import streamlit as st
except ImportError as exc:  # pragma: no cover
    raise SystemExit("Streamlit is not installed. Install requirements.txt and run: streamlit run app/ui/streamlit_app.py") from exc

from app.config import Settings
from app.detectors.context_registry import DetectorPipeline
from app.processors.text_processor import TextProcessor
from app.pseudonymization.mapper import Pseudonymizer
from app.processors.docx_processor import DocxProcessor
from app.evaluation.evaluator import Evaluator
from app.main import make_benchmark


st.set_page_config(page_title="PII Redaction Studio", page_icon="🛡️", layout="wide")
st.markdown("""
<style>
.block-container {max-width: 1200px; padding-top: 2rem;}
.metric-card {padding: 1rem 1.2rem; border: 1px solid rgba(127,127,127,.25); border-radius: 14px; background: rgba(127,127,127,.06);}
.small {font-size:.88rem; opacity:.75;}
</style>
""", unsafe_allow_html=True)

st.title("PII Redaction Studio")
st.caption("Local DOCX PII detection, deterministic pseudonymization, and evaluation")

with st.sidebar:
    st.header("Settings")
    threshold=st.slider("Confidence threshold", 0.50, 0.99, 0.80, 0.01)
    seed=st.number_input("Deterministic seed", min_value=1, value=20250925, step=1)
    image_ocr=st.checkbox("Process embedded images with OCR", value=True)
    debug=st.checkbox("Developer debug mode", value=False, help="May expose entity values. Keep disabled for normal use.")

uploaded=st.file_uploader("Upload a DOCX document", type=["docx"])

if uploaded:
    with tempfile.TemporaryDirectory() as td:
        in_path=Path(td)/uploaded.name
        out_path=Path(td)/"redacted_output.docx"
        in_path.write_bytes(uploaded.getbuffer())
        if st.button("Analyze and redact", type="primary"):
            processor=DocxProcessor(TextProcessor(DetectorPipeline(), threshold), Pseudonymizer(seed))
            with st.status("Processing document...", expanded=False) as status:
                stats=processor.process(in_path, out_path)
                status.update(label="Completed", state="complete")
            st.subheader("Overview")
            cols=st.columns(4)
            for c,label,value in zip(cols,["Paragraphs","Tables","Images","Entities"],[stats.paragraphs,stats.tables,stats.embedded_images,len(stats.entities)]):
                with c:
                    st.markdown(f'<div class="metric-card"><b>{label}</b><h2>{value}</h2></div>', unsafe_allow_html=True)
            st.divider()
            tabs=st.tabs(["Detection Results","Redacted Output","Evaluation","Logs"])
            with tabs[0]:
                st.dataframe([{"Type":k,"Count":v} for k,v in sorted(stats.counts.items())], use_container_width=True)
                if debug:
                    st.dataframe([{"type":e.type.value,"value":e.value,"confidence":e.confidence,"source":e.source,"location":e.context.location} for e in stats.entities], use_container_width=True)
                else:
                    st.info("Entity values are hidden. Enable Developer debug mode only in a trusted local session.")
            with tabs[1]:
                st.download_button("Download redacted DOCX", data=out_path.read_bytes(), file_name="redacted_output.docx", mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
                st.success("A redacted DOCX was generated locally.")
                try:
                    from docx import Document
                    preview=Document(str(out_path))
                    sample=[p.text for p in preview.paragraphs if p.text.strip()][:8]
                    st.subheader("Redacted text preview")
                    st.code("\n".join(sample), language="text")
                except Exception as exc:
                    st.warning(f"Preview unavailable: {exc}")
            with tabs[2]:
                if st.button("Run benchmark evaluation"):
                    with st.spinner("Evaluating detector benchmark..."):
                        eval_result=Evaluator(TextProcessor(DetectorPipeline(), threshold)).evaluate_cases(make_benchmark())
                    st.subheader("Overall")
                    eco=st.columns(4)
                    overall=eval_result["overall"]
                    for c,label,key in zip(eco,["Precision","Recall","F1","Accuracy"],["precision","recall","f1","character_accuracy"]):
                        with c:
                            st.metric(label, f"{overall[key]*100:.1f}%")
                    rows=[]
                    for typ,m in sorted(eval_result["per_category"].items()):
                        rows.append({"Type":typ,"TP":m["tp"],"FP":m["fp"],"FN":m["fn"],"Precision":m["precision"],"Recall":m["recall"],"F1":m["f1"]})
                    st.dataframe(rows, use_container_width=True)
                else:
                    st.info("Run the benchmark to view measured precision, recall, F1, and character-level accuracy.")
            with tabs[3]:
                st.write({"paragraphs":stats.paragraphs,"tables":stats.tables,"table_cells":stats.table_cells,"headers":stats.headers,"footers":stats.footers,"embedded_images":stats.embedded_images,"ocr_images":stats.ocr_images,"redacted_images":stats.redacted_images})
else:
    st.info("Upload a DOCX to begin.")
