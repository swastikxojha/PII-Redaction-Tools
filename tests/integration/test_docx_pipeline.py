from docx import Document
from app.detectors.context_registry import DetectorPipeline
from app.processors.text_processor import TextProcessor
from app.pseudonymization.mapper import Pseudonymizer
from app.processors.docx_processor import DocxProcessor


def test_docx_end_to_end(tmp_path):
    src=tmp_path/'input.docx'; out=tmp_path/'out.docx'
    d=Document(); p=d.add_paragraph('Contact: John Smith, Email john.smith@example.com, Phone +91 9876543210'); d.save(src)
    stats=DocxProcessor(TextProcessor(DetectorPipeline()),Pseudonymizer(123)).process(src,out)
    assert out.exists() and stats.counts.get('EMAIL')==1
    result=Document(out)
    txt=' '.join(p.text for p in result.paragraphs)
    assert 'john.smith@example.com' not in txt
