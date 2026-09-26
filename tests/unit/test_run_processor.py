from docx import Document
from app.processors.run_processor import apply_replacements_to_paragraph
from app.models.entity import PIIEntity, PIIType


def test_replacement_across_runs(tmp_path):
    d=Document(); p=d.add_paragraph()
    r1=p.add_run('Email: john@')
    r2=p.add_run('example.com')
    entity=PIIEntity(PIIType.EMAIL,'john@example.com',7,23,.99,'test')
    apply_replacements_to_paragraph(p,[(entity,'jane@example.com')])
    assert p.text == 'Email: jane@example.com'

def test_multiple_replacements_do_not_corrupt_neighboring_spans():
    d=Document(); p=d.add_paragraph()
    p.add_run('John Smith ')
    p.add_run('john@example.com ')
    p.add_run('+91 9876543210')
    e1=PIIEntity(PIIType.PERSON,'John Smith',0,10,.99,'test')
    e2=PIIEntity(PIIType.EMAIL,'john@example.com',11,27,.99,'test')
    e3=PIIEntity(PIIType.PHONE,'+91 9876543210',28,42,.99,'test')
    apply_replacements_to_paragraph(p,[(e1,'Jane Doe'),(e2,'jane@example.com'),(e3,'+91 9000000000')])
    assert p.text == 'Jane Doe jane@example.com +91 9000000000'
